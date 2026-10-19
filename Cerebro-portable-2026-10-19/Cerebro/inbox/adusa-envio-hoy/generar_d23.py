#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el D23 (jueves 17 sep 2026) a partir del D22.

El miercoles 16 no salio tracker, asi que la letra no avanzo ese dia.
D19 el 10, D20 el 11, D21 el 14, D22 el 15, D23 el 17.

Cambios de hoy, todos verificados en Jira en vivo:
  RCS-28910  Closed, Resolved, cerrado el 9/17
  RCS-28944  Closed, Resolved, cerrado el 9/16, asignado a Vlad Bogus
  RCS-28911  Completed, Resolved, cerrado el 9/16
  RCS-28824  vuelve a Work in progress, asignado a Sebastian Rueda
  RCS-28281  seguimiento real del 9/17, sigue Awaiting Ahold input
  RCS-28683  solo texto, el item de RCS-28910 continua aqui. Sin cambio de fecha.

nueva() se conserva para dias con tickets nuevos. Hoy no se usa.
"""

import sys, os, re, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    AQUI, "Appriss ADUSA-Support-Tracker - 09152026 D22.xlsx")
DESTINO = os.path.join(AQUI, "Appriss ADUSA-Support-Tracker - 09172026 D23.xlsx")
if not os.path.exists(ORIGEN):
    sys.exit("No encuentro el D22 en:\n  %s" % ORIGEN)

FILL = {"Resolved": "FF7DE2DD", "Active with Appriss": "FFFBC521",
        "Awaiting Ahold input": "FF7C6BE8", "Pending Ahold verification": "FF528DF5"}
BLANCO, BANDA = "FFFFFFFF", "FFF4F6F9"
HOY = datetime.datetime(2026, 9, 17)
AYER = datetime.datetime(2026, 9, 16)

wb = openpyxl.load_workbook(ORIGEN)
ws = wb["Support tracker"]


def need(clave):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == clave:
            return r
    sys.exit("No encontre %s. Revisa que el archivo sea el D22." % clave)


def categoria(r, cat):
    c = ws.cell(row=r, column=4); c.value = cat
    v = c.fill
    c.fill = PatternFill(fill_type=v.fill_type, fgColor=FILL[cat], bgColor=v.bgColor)


def anexa(r, linea):
    c = ws.cell(row=r, column=10)
    c.value = (c.value or "") + "\n" + linea


def bloque(nombre):
    hdr = None
    for r in range(23, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == nombre:
            hdr = r; break
    if hdr is None:
        sys.exit("No encontre el bloque '%s'." % nombre)
    ultima = hdr
    for r in range(hdr + 1, ws.max_row + 2):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().startswith("RCS-"):
            ultima = r
        elif v:
            break
    return ultima


def nueva(nombre_tema, vals, alto):
    plantilla = bloque(nombre_tema)
    dest = plantilla + 1
    guardadas = [str(m) for m in list(ws.merged_cells.ranges) if m.min_row >= dest]
    for s in guardadas:
        ws.unmerge_cells(s)
    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if d.height is not None and r >= dest}
    ws.insert_rows(dest, amount=1)
    for s in guardadas:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", s)
        ws.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1,
                                      m.group(3), int(m.group(4)) + 1))
    for r in sorted(alturas, reverse=True):
        ws.row_dimensions[r + 1].height = alturas[r]
    ws.row_dimensions[dest].height = alto
    for c in range(1, 11):
        ws.cell(row=dest, column=c)._style = copy(ws.cell(row=plantilla, column=c)._style)
    fg = ws.cell(row=plantilla, column=1).fill.fgColor
    previa = fg.rgb if fg.type == "rgb" else BLANCO
    banda = BLANCO if previa == BANDA else BANDA
    for c in range(1, 11):
        ws.cell(row=dest, column=c).fill = PatternFill(
            fill_type="solid", fgColor=FILL[vals[3]] if c == 4 else banda)
    for i, v in enumerate(vals, start=1):
        ws.cell(row=dest, column=i).value = v
    for c in (6, 7):
        ws.cell(row=dest, column=c).number_format = \
            ws.cell(row=plantilla, column=c).number_format


ws["A2"] = ("Status as of Thursday, September 17, 2026. "
            "Ticket states verified directly against Jira.")

# ---------- RCS-28910, cerrado ----------
r = need("RCS-28910")
ws.cell(row=r, column=3).value = "Closed"
categoria(r, "Resolved")
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "The Elevator Entrapment case type now carries the side bar and the case information bar that "
    "the other case types show, and Update Status from the workbench now offers both Case "
    "Completed and Closed. Leigh Ann verified both on 9/16 and agreed on 9/17 that the ticket "
    "could close. The error that appears when adding a person to a case is tracked on RCS-28683 "
    "and continues there.")
anexa(r, "9/16/2026 - Side bar and case information bar added to the workbench for this case type, "
         "and Update Status now allows the case to be completed as well as closed.")
anexa(r, "9/17/2026 - Ahold confirmed the case type behaves correctly and agreed to close. The "
         "error when adding a person to a case moves to RCS-28683.")
ws.row_dimensions[r].height = 160

# ---------- RCS-28944, cerrado ----------
r = need("RCS-28944")
ws.cell(row=r, column=3).value = "Closed"
categoria(r, "Resolved")
ws.cell(row=r, column=5).value = "Vlad Bogus"
ws.cell(row=r, column=7).value = AYER
ws.cell(row=r, column=9).value = (
    "Appriss walked Ahold through rebuilding the mapping on the report column so it reads Active "
    "when the value is zero and Inactive otherwise. Jim confirmed on 9/16 that the report returns "
    "the expected text, and the ticket closed the same day.")
anexa(r, "9/16/2026 - Appriss provided the steps to rebuild the mapping on the report column.")
anexa(r, "9/16/2026 - That worked, thank you.")
ws.row_dimensions[r].height = 135

# ---------- RCS-28911, cerrado ----------
r = need("RCS-28911")
ws.cell(row=r, column=3).value = "Completed"
categoria(r, "Resolved")
ws.cell(row=r, column=7).value = AYER
ws.cell(row=r, column=9).value = (
    "The store and location update was not part of the load that ran on the morning of 9/14. Ahold "
    "confirmed on 9/16 that the file went out late because of server issues on its side, and that "
    "the file was picked up correctly on the next scheduled run. Jim and Leigh Ann agreed the "
    "ticket could close, and it closed the same day.")
anexa(r, "9/16/2026 - The file was sent late as a result of server issues on our side. That file "
         "or the new file sent Tuesday was processed Tuesday morning correctly, so this can be "
         "closed.")
ws.row_dimensions[r].height = 150

# ---------- RCS-28824, dos puntos abiertos, pasa al equipo de datos ----------
r = need("RCS-28824")
ws.cell(row=r, column=3).value = "Work in progress"
ws.cell(row=r, column=5).value = "Sebastian Rueda"
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "Appriss corrected the store profile and Becky confirmed the Week, Month, Quarter and Year "
    "tabs now return charts. Two points stay open and both look like missing data rather than the "
    "pages themselves. The cashier profile page returns no charts, and the area map has no store "
    "coordinates to draw on. Ahold confirmed on 9/17 that it sends a store and location file and "
    "does not send an employee file, and that what it needs on the cashier page is for the charts "
    "to populate rather than employee detail. The ticket sits with the Appriss data team to "
    "confirm what arrives on each feed and what is needed to make those charts return data.")
anexa(r, "9/16/2026 - The empty cashier profile charts and the blank area map both point to "
         "missing data rather than to the pages. Referred to the data team.")
anexa(r, "9/17/2026 - We do send a location and store file, we do not send an employee file. What "
         "we need on the cashier profile page is for the charts to populate with refunds, voids "
         "and coupon history, not employee detail.")
ws.row_dimensions[r].height = 165

# ---------- RCS-28281, seguimiento real del 9/17 ----------
r = need("RCS-28281")
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "Appriss posted the full count per trading day on 8/27. Ahold reported on 8/31 that a good "
    "deal of recovery had occurred, with two gaps left on its view, the 9 PM to 10 PM hour on "
    "8/12 and a handful of trading days sitting 100,000 to 200,000 transactions below average. "
    "Appriss asked Rodney on 9/7 to have Becky run the same 10 minute check she used on 8/20 "
    "against 8/11, 8/12, 8/19, 8/20, 8/22 and 8/25, so Appriss can confirm which batches still "
    "need resubmission. Appriss followed up on 9/17 and also asked whether service has returned "
    "to normal on Ahold's side.")
anexa(r, "9/17/2026 - Following up on the ten minute check for the flagged dates. Please let us "
         "know if there are any updates on this incident or if service has returned to normal.")
ws.row_dimensions[r].height = 105

# ---------- RCS-28683, solo texto. No se toca la fecha, no hubo actividad en su hilo ----------
r = need("RCS-28683")
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value.rstrip() +
    " Leigh Ann linked this on 9/16 to the same error seen on the Elevator Entrapment case type, "
    "and that item now continues here.")
ws.row_dimensions[r].height = 130

# ---------- resumen ----------
def fres(nombre):
    for r in range(7, 18):
        if ws.cell(row=r, column=1).value == nombre:
            return r
    sys.exit("No encontre '%s' en el resumen." % nombre)

r = fres("Case entry and assignment")
ws.cell(row=r, column=4).value = 9; ws.cell(row=r, column=7).value = 2
r = fres("Reporting and exports")
ws.cell(row=r, column=4).value = 5; ws.cell(row=r, column=7).value = 1
r = fres("Store file and location data")
ws.cell(row=r, column=4).value = 3; ws.cell(row=r, column=6).value = 0

wb.save(DESTINO)

# ---------- verificacion ----------
ws2 = openpyxl.load_workbook(DESTINO)["Support tracker"]
print("\nGuardado:", os.path.basename(DESTINO)); print(ws2["A2"].value, "\n")
tot = [0, 0, 0, 0]
for r in range(7, 18):
    for i, c in enumerate((4, 5, 6, 7)):
        tot[i] += ws2.cell(row=r, column=c).value or 0
print("resueltos %d | pend verif %d | esperando Ahold %d | activos %d | total %d"
      % (tot[0], tot[1], tot[2], tot[3], sum(tot)))
esp = {"RCS-28693": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28667": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28683": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28849": ("Accepted into Development", "Active with Appriss"),
       "RCS-28824": ("Work in progress",          "Active with Appriss"),
       "RCS-28589": ("Accepted into Development", "Active with Appriss"),
       "RCS-28281": ("Under Customer Review",     "Awaiting Ahold input")}
ok, vistos, claves = True, set(), []
print("\nAbiertos:")
for r in range(23, ws2.max_row + 1):
    k = ws2.cell(row=r, column=1).value
    if not (k and str(k).startswith("RCS-")):
        continue
    claves.append(k)
    cat = ws2.cell(row=r, column=4).value
    if cat == "Resolved":
        continue
    est = ws2.cell(row=r, column=3).value
    vistos.add(k)
    mal = k in esp and (est, cat) != esp[k]
    ok = ok and not mal
    print("  %s fila %-4d %-11s %-26s %s" % ("!!" if mal else "  ", r, k, est, cat))
if set(esp) - vistos:
    ok = False; print("\nFaltan:", ", ".join(sorted(set(esp) - vistos)))
if vistos - set(esp):
    ok = False; print("\nAbiertos inesperados:", ", ".join(sorted(vistos - set(esp))))
if sum(tot) != 55 or tot[0] != 48:
    ok = False; print("\nTotales mal, deberian ser 55 y 48 resueltos.")
print("\nCerrados hoy y ayer:")
for k in ("RCS-28910", "RCS-28944", "RCS-28911"):
    for r in range(23, ws2.max_row + 1):
        if str(ws2.cell(row=r, column=1).value).strip() == k:
            c4 = ws2.cell(row=r, column=4)
            print("  fila %-4d %-11s %-11s %-9s fill=%s" % (
                r, k, ws2.cell(row=r, column=3).value, c4.value, c4.fill.fgColor.rgb))
            if c4.value != "Resolved" or c4.fill.fgColor.rgb != FILL["Resolved"]:
                ok = False
            break
d = sorted({k for k in claves if claves.count(k) > 1})
if d:
    ok = False; print("\nDuplicados:", ", ".join(d))
print("\n" + ("Todo cuadra." if ok else "Revisa lo marcado."))
