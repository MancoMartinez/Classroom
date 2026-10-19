#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el D24 (viernes 18 sep 2026) a partir del D23.

D19 el 10, D20 el 11, D21 el 14, D22 el 15, D23 el jueves 17, D24 el viernes 18.

Dia tranquilo. Verificado en Jira en vivo:
  RCS-29000  nuevo, abierto por Becky LaPierre el 9/18, bloque Reporting and exports
  Los siete abiertos del D23 sin cambio de estado ni de asignado desde el D23.
"""

import sys, os, re, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    AQUI, "Appriss ADUSA-Support-Tracker - 09172026 D23.xlsx")
DESTINO = os.path.join(AQUI, "Appriss ADUSA-Support-Tracker - 09182026 D24.xlsx")
if not os.path.exists(ORIGEN):
    sys.exit("No encuentro el D23 en:\n  %s" % ORIGEN)

FILL = {"Resolved": "FF7DE2DD", "Active with Appriss": "FFFBC521",
        "Awaiting Ahold input": "FF7C6BE8", "Pending Ahold verification": "FF528DF5"}
BLANCO, BANDA, HOY = "FFFFFFFF", "FFF4F6F9", datetime.datetime(2026, 9, 18)

wb = openpyxl.load_workbook(ORIGEN)
ws = wb["Support tracker"]


def need(clave):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == clave:
            return r
    sys.exit("No encontre %s. Revisa que el archivo sea el D23." % clave)


def categoria(r, cat):
    c = ws.cell(row=r, column=4); c.value = cat
    v = c.fill
    c.fill = PatternFill(fill_type=v.fill_type, fgColor=FILL[cat], bgColor=v.bgColor)


def anexa(r, linea):
    c = ws.cell(row=r, column=10)
    c.value = (c.value or "") + "\n" + linea


def bloque(nombre):
    hdr = None
    for r in range(23, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == nombre:
            hdr = r; break
    if hdr is None:
        sys.exit("No encontre el bloque '%s'." % nombre)
    ultima = hdr
    for r in range(hdr + 1, ws.max_row + 2):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().startswith("RCS-"):
            ultima = r
        elif v:
            break
    return ultima


def nueva(nombre_tema, vals, alto):
    plantilla = bloque(nombre_tema)
    dest = plantilla + 1
    guardadas = [str(m) for m in list(ws.merged_cells.ranges) if m.min_row >= dest]
    for s in guardadas:
        ws.unmerge_cells(s)
    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if d.height is not None and r >= dest}
    ws.insert_rows(dest, amount=1)
    for s in guardadas:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", s)
        ws.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1,
                                      m.group(3), int(m.group(4)) + 1))
    for r in sorted(alturas, reverse=True):
        ws.row_dimensions[r + 1].height = alturas[r]
    ws.row_dimensions[dest].height = alto
    for c in range(1, 11):
        ws.cell(row=dest, column=c)._style = copy(ws.cell(row=plantilla, column=c)._style)
    fg = ws.cell(row=plantilla, column=1).fill.fgColor
    previa = fg.rgb if fg.type == "rgb" else BLANCO
    banda = BLANCO if previa == BANDA else BANDA
    for c in range(1, 11):
        ws.cell(row=dest, column=c).fill = PatternFill(
            fill_type="solid", fgColor=FILL[vals[3]] if c == 4 else banda)
    for i, v in enumerate(vals, start=1):
        ws.cell(row=dest, column=i).value = v
    for c in (6, 7):
        ws.cell(row=dest, column=c).number_format = \
            ws.cell(row=plantilla, column=c).number_format


ws["A2"] = ("Status as of Friday, September 18, 2026. "
            "Ticket states verified directly against Jira.")

# ---------- RCS-29000, nuevo ----------
nueva("Reporting and exports", [
    "RCS-29000",
    "Percent columns round inconsistently on a saved report",
    "Under Support Review", "Active with Appriss", "Sebastian Rueda", HOY, HOY,
    "Percent columns on a front end report showed two decimal places until new columns were added "
    "to it. Since then one or two of those columns inconsistently return more than two decimals, "
    "and the column that behaves that way changes between runs of the same saved search, for the "
    "same store and for different stores. Ahold asks whether the rounding can be set on the fields "
    "or whether the behaviour comes from the application. Raised 9/18 with screenshots.",
    "Open with Appriss. Raised on 9/18 and under support review. Appriss took ownership the same "
    "day and is reviewing how the rounding is applied on those display fields before coming back "
    "with an answer.",
    "Ticket opened 9/18/2026\n9/18/2026 - Percent columns used to show two decimal places. After "
    "new columns were added, one or two of them inconsistently show more than two, and the "
    "affected column changes from run to run. Is this the application, or is there a way to set "
    "the rounding on these fields?",
], 150)

# ---------- resumen ----------
def fres(nombre):
    for r in range(7, 18):
        if ws.cell(row=r, column=1).value == nombre:
            return r
    sys.exit("No encontre '%s' en el resumen." % nombre)

ws.cell(row=fres("Reporting and exports"), column=7).value = 2

wb.save(DESTINO)

# ---------- verificacion ----------
ws2 = openpyxl.load_workbook(DESTINO)["Support tracker"]
print("\nGuardado:", os.path.basename(DESTINO)); print(ws2["A2"].value, "\n")
tot = [0, 0, 0, 0]
for r in range(7, 18):
    for i, c in enumerate((4, 5, 6, 7)):
        tot[i] += ws2.cell(row=r, column=c).value or 0
print("resueltos %d | pend verif %d | esperando Ahold %d | activos %d | total %d"
      % (tot[0], tot[1], tot[2], tot[3], sum(tot)))
esp = {"RCS-28693": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28667": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28683": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28849": ("Accepted into Development", "Active with Appriss"),
       "RCS-28824": ("Work in progress",          "Active with Appriss"),
       "RCS-29000": ("Under Support Review",      "Active with Appriss"),
       "RCS-28589": ("Accepted into Development", "Active with Appriss"),
       "RCS-28281": ("Under Customer Review",     "Awaiting Ahold input")}
ok, vistos, claves = True, set(), []
print("\nAbiertos:")
for r in range(23, ws2.max_row + 1):
    k = ws2.cell(row=r, column=1).value
    if not (k and str(k).startswith("RCS-")):
        continue
    claves.append(k)
    cat = ws2.cell(row=r, column=4).value
    if cat == "Resolved":
        continue
    est = ws2.cell(row=r, column=3).value
    vistos.add(k)
    mal = k in esp and (est, cat) != esp[k]
    ok = ok and not mal
    print("  %s fila %-4d %-11s %-26s %s" % ("!!" if mal else "  ", r, k, est, cat))
if set(esp) - vistos:
    ok = False; print("\nFaltan:", ", ".join(sorted(set(esp) - vistos)))
if vistos - set(esp):
    ok = False; print("\nAbiertos inesperados:", ", ".join(sorted(vistos - set(esp))))
if sum(tot) != 56 or tot[0] != 48:
    ok = False; print("\nTotales mal, deberian ser 56 y 48 resueltos.")
d = sorted({k for k in claves if claves.count(k) > 1})
if d:
    ok = False; print("\nDuplicados:", ", ".join(d))
print("\n" + ("Todo cuadra." if ok else "Revisa lo marcado."))
