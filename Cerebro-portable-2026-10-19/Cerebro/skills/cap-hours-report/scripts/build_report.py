"""
Capitalizable Hours Excel report builder — Appriss Retail Engineering.
Usage:
    python build_report.py --input classified.json --output "Capitalizable Hours - April 2026.xlsx"
    python build_report.py classified.json "Capitalizable Hours - April 2026.xlsx"   # positional also accepted
"""
import sys
import json
import argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Accept either named args or positional for convenience
if "--input" in sys.argv or "--output" in sys.argv:
    p = argparse.ArgumentParser()
    p.add_argument("--input",  required=True)
    p.add_argument("--output", required=True)
    a = p.parse_args()
    json_path = a.input
    out_path  = a.output
else:
    json_path = sys.argv[1]
    out_path  = sys.argv[2]

with open(json_path) as f:
    data = json.load(f)

results        = data["results"]
summary        = data["summary"]
WORKING_DAYS   = data["working_days"]
PRODUCTIVE_HRS = data["productive_hrs"]
MONTH          = data["month"]

# ─── Colour palette ──────────────────────────────────────────────────────────
NAVY      = "002060"
WHITE     = "FFFFFF"
LIGHT_GREY = "F2F2F2"
GREEN_FILL = "E2EFDA"
RED_FILL   = "FCE4D6"   # data-quality warning background
RED_FONT   = "9C0006"   # data-quality warning text

LOW_POINT_THRESHOLD = 10  # engineers below this with any cap pts get flagged

def hdr_font(bold=True, color=WHITE, sz=11):
    return Font(name="Arial", bold=bold, color=color, size=sz)

def cell_font(bold=False, color="000000", sz=10):
    return Font(name="Arial", bold=bold, color=color, size=sz)

def fill(hex_col):
    return PatternFill("solid", fgColor=hex_col)

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# ════════════════════════════════════════════════════════════════════════════
#  TAB 1 – SUMMARY
# ════════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Summary"
ws.sheet_view.showGridLines = False

# Title row
ws.merge_cells("A1:N1")
t = ws["A1"]
t.value = f"Capitalizable Hours Report — {MONTH}"
t.font  = Font(name="Arial", bold=True, color=WHITE, size=14)
t.fill  = fill(NAVY)
t.alignment = center()
ws.row_dimensions[1].height = 28

# Methodology box
method_text = (
    f"Points → Hours Conversion (per engineer):\n"
    f"  • Available hours = {WORKING_DAYS} working days × 6 productive hrs/day = {PRODUCTIVE_HRS} hrs/engineer\n"
    f"  • Hours per point = {PRODUCTIVE_HRS} ÷ engineer's total story points\n"
    f"  • Capitalizable hours = capitalizable points × hours per point\n"
    f"  • Engineers with 0 total points are shown with 0 hrs (no conversion possible)"
)
ws.merge_cells("A2:N4")
m = ws["A2"]
m.value = method_text
m.font  = Font(name="Arial", size=9, italic=True, color="404040")
m.fill  = fill("EEF3F9")
m.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 15
ws.row_dimensions[3].height = 15
ws.row_dimensions[4].height = 15

# Column headers
HDR_ROW = 6
headers = [
    "Engineer", "Total Items", "Total\nStory Points",
    "Capitalizable\nItems", "Capitalizable\nPoints",
    "Hrs / Point", "Available\nHours",
    "Capitalizable\nHours", "Non-Cap\nHours", "% Capitalizable",
    "Engage\nCap Hrs", "Secure\nCap Hrs", "Other\nCap Hrs",
    "Data Quality Flag"
]
ws.row_dimensions[HDR_ROW].height = 36
for col, h in enumerate(headers, 1):
    c = ws.cell(row=HDR_ROW, column=col, value=h)
    c.font      = hdr_font()
    c.fill      = fill(NAVY)
    c.alignment = center()
    c.border    = thin_border()

# Data rows
DATA_START = HDR_ROW + 1
for i, row in enumerate(summary):
    r = DATA_START + i
    ws.row_dimensions[r].height = 16

    # Flag: low total points with some capitalizable → potential over-reporting
    is_flagged = (row["total_pts"] < LOW_POINT_THRESHOLD and row["cap_pts"] > 0)
    if is_flagged:
        bg = fill(RED_FILL)
    else:
        bg = fill(LIGHT_GREY) if i % 2 == 0 else fill(WHITE)

    flag_text = (
        f"⚠ Only {row['total_pts']} pts recorded — may not reflect full workload; "
        "capitalizable hours could be overstated."
        if is_flagged else ""
    )

    values = [
        row["engineer"],
        row["item_count"],
        row["total_pts"],
        row["cap_item_count"],
        row["cap_pts"],
        round(row["hrs_per_pt"], 2) if row["total_pts"] > 0 else "—",
        PRODUCTIVE_HRS,
        round(row["cap_hrs"], 1),
        round(row["total_pts"] * row["hrs_per_pt"] - row["cap_hrs"], 1) if row["total_pts"] > 0 else 0,
        row["cap_pct"] / 100 if row["total_pts"] > 0 else None,
        round(row.get("engage_cap_hrs", 0.0), 1),
        round(row.get("secure_cap_hrs", 0.0), 1),
        round(row.get("other_cap_hrs",  0.0), 1),
        flag_text,
    ]
    fmts   = [None, "#,##0", "#,##0", "#,##0", "#,##0",
              "#,##0.00", "#,##0", "#,##0.0", "#,##0.0", "0.0%",
              "#,##0.0", "#,##0.0", "#,##0.0", None]
    aligns = ["left"] + ["center"] * 12 + ["left"]

    for col, (val, fmt, aln) in enumerate(zip(values, fmts, aligns), 1):
        c = ws.cell(row=r, column=col, value=val)
        c.border    = thin_border()
        c.alignment = Alignment(horizontal=aln, vertical="center", wrap_text=(col == 14))
        if fmt and val not in (None, "—"):
            c.number_format = fmt
        if is_flagged:
            c.fill = fill(RED_FILL)
            c.font = cell_font(bold=(col == 14), color=RED_FONT)
        else:
            c.fill = bg
            c.font = cell_font()

    # Green highlight for cap hours / cap% / product cols (only on non-flagged rows)
    if not is_flagged:
        cap_cell = ws.cell(row=r, column=8)
        if row["cap_hrs"] > 0:
            cap_cell.fill = fill(GREEN_FILL)
            cap_cell.font = cell_font(bold=True, color="375623")

        pct_cell = ws.cell(row=r, column=10)
        if row["total_pts"] > 0 and row["cap_pct"] > 0:
            pct_cell.fill = fill("C6EFCE")
            pct_cell.font = cell_font(bold=True, color="276221")

        # Soft teal tint for product breakdown columns (11=Engage, 12=Secure, 13=Other)
        TEAL_FILL = "DDEEFF"
        for prod_col in (11, 12, 13):
            pc = ws.cell(row=r, column=prod_col)
            if pc.value and pc.value > 0:
                pc.fill = fill(TEAL_FILL)
                pc.font = cell_font(bold=True, color="1F497D")

# Totals row
TOTAL_ROW = DATA_START + len(summary)
ws.row_dimensions[TOTAL_ROW].height = 18
tc = ws.cell(row=TOTAL_ROW, column=1, value="TOTALS / TEAM AVERAGE")
tc.font      = hdr_font(sz=10)
tc.fill      = fill(NAVY)
tc.alignment = Alignment(horizontal="left", vertical="center")
tc.border    = thin_border()
for _ec in range(2, 7):   # cols 2-6 are blank navy filler before the numeric totals start at col 7
    _c = ws.cell(row=TOTAL_ROW, column=_ec, value="")
    _c.fill = fill(NAVY); _c.font = hdr_font(sz=10); _c.border = thin_border()

total_cap_hrs = sum(r["cap_hrs"] for r in summary)
total_non_cap = sum((r["total_pts"] * r["hrs_per_pt"] - r["cap_hrs"])
                    for r in summary if r["total_pts"] > 0)
total_pts_all = sum(r["total_pts"] for r in summary)
total_cap_pts = sum(r["cap_pts"] for r in summary)

flagged_count = sum(
    1 for row in summary if row["total_pts"] < LOW_POINT_THRESHOLD and row["cap_pts"] > 0
)
flag_summary = (
    f"{flagged_count} engineer(s) flagged — review capitalizable hrs above"
    if flagged_count > 0 else ""
)
total_engage = sum(r.get("engage_cap_hrs", 0.0) for r in summary)
total_secure = sum(r.get("secure_cap_hrs", 0.0) for r in summary)
total_other  = sum(r.get("other_cap_hrs",  0.0) for r in summary)

totals = ["", "", "", "", "", PRODUCTIVE_HRS,
          round(total_cap_hrs, 1), round(total_non_cap, 1),
          round(total_cap_hrs / (total_cap_hrs + total_non_cap) * 100, 1) / 100 if (total_cap_hrs + total_non_cap) > 0 else None,
          round(total_engage, 1), round(total_secure, 1), round(total_other, 1),
          flag_summary]
fmts2  = [None, None, None, None, None, "#,##0", "#,##0.0", "#,##0.0", "0.0%",
          "#,##0.0", "#,##0.0", "#,##0.0", None]
for col, (val, fmt) in enumerate(zip(totals, fmts2), 2):
    c = ws.cell(row=TOTAL_ROW, column=col, value=val)
    c.font   = hdr_font(sz=10)
    c.fill   = fill(NAVY)
    c.border = thin_border()
    c.alignment = center()
    if fmt and val not in (None, ""):
        c.number_format = fmt

col_widths = [28, 10, 12, 12, 12, 10, 10, 14, 12, 14, 12, 12, 12, 52]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = f"A{DATA_START}"

# ════════════════════════════════════════════════════════════════════════════
#  TAB 2 – DETAIL
# ════════════════════════════════════════════════════════════════════════════
wd = wb.create_sheet("Work Item Detail")
wd.sheet_view.showGridLines = False

wd.merge_cells("A1:N1")
dt = wd["A1"]
dt.value = f"Work Item Detail — {MONTH}  ({len(results)} items from ENG & DATASCI projects)"
dt.font  = Font(name="Arial", bold=True, color=WHITE, size=12)
dt.fill  = fill(NAVY)
dt.alignment = center()
wd.row_dimensions[1].height = 24

# Legend row: explain AI-assisted amber colouring
wd.merge_cells("A2:N2")
leg = wd["A2"]
leg.value = (
    "Status colours: Green = Capitalizable (rule-based) | "
    "Amber = Capitalizable or Non-Capitalizable (AI-Assisted — review recommended) | "
    "Red = Non-Capitalizable (rule-based)"
)
leg.font      = Font(name="Arial", size=9, italic=True, color="404040")
leg.fill      = fill("FFFBF0")
leg.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
wd.row_dimensions[2].height = 14

D_HDR = 3
det_headers = [
    "Jira Key", "Assignee", "Type", "Story Points",
    "Work Category", "Epic", "Epic Summary", "Initiative", "Initiative Name",
    "Product Category", "Resolution Date", "Status", "Reasoning", "Summary"
]
wd.row_dimensions[D_HDR].height = 30
for col, h in enumerate(det_headers, 1):
    c = wd.cell(row=D_HDR, column=col, value=h)
    c.font      = hdr_font()
    c.fill      = fill(NAVY)
    c.alignment = center()
    c.border    = thin_border()

D_DATA = D_HDR + 1
sorted_results = sorted(results, key=lambda x: (x["assignee"], x["key"]))

for i, item in enumerate(sorted_results):
    r = D_DATA + i
    wd.row_dimensions[r].height = 14
    bg = fill(LIGHT_GREY) if i % 2 == 0 else fill(WHITE)
    is_cap = item["cap_status"] in ("Capitalizable", "Capitalizable (AI-Assisted)")
    is_ai  = "(AI-Assisted)" in item.get("cap_status", "")

    pts      = item.get("points")
    res_date = (item.get("resolutiondate") or "")[:10]

    # Flatten any dict fields (e.g. work_category may be {"value": "Bug", ...})
    def _str(v):
        if isinstance(v, dict):
            return v.get("value", str(v))
        return v if v is not None else ""

    prod_cat = item.get("product_category", "Other")

    row_vals = [
        item["key"],
        item["assignee"],
        item.get("type", ""),
        pts if pts is not None else "",
        _str(item.get("work_category")),
        item.get("resolved_epic", ""),
        item.get("resolved_epic_summary", ""),
        item.get("resolved_initiative", ""),
        item.get("resolved_initiative_name", ""),
        prod_cat,
        res_date,
        item["cap_status"],
        item["reasoning"],
        item["summary"],
    ]

    for col, val in enumerate(row_vals, 1):
        c = wd.cell(row=r, column=col, value=val)
        c.font      = cell_font()
        c.border    = thin_border()
        c.alignment = Alignment(horizontal="left", vertical="center",
                                wrap_text=(col in (13, 14)))
        if col == 12:   # "Status" column — green/red/amber by cap status
            if is_ai:
                # AI-assisted: amber background to prompt review
                c.fill = fill("FFF2CC")
                c.font = cell_font(bold=True, color="7F6000")
            elif is_cap:
                c.fill = fill("C6EFCE")
                c.font = cell_font(bold=True, color="276221")
            else:
                c.fill = fill("FCE4D6")
                c.font = cell_font(bold=True, color="9C0006")
        elif col == 10:  # "Product Category" — subtle teal tint
            TEAL_FILL = "DDEEFF"
            c.fill = fill(TEAL_FILL)
            c.font = cell_font(bold=True, color="1F497D")
        else:
            c.fill = fill(GREEN_FILL) if is_cap else bg

    key_cell = wd.cell(row=r, column=1)
    key_cell.hyperlink = f"https://trebuco.atlassian.net/browse/{item['key']}"
    key_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

det_col_widths = [14, 24, 10, 12, 20, 14, 34, 14, 32, 14, 13, 16, 50, 60]
for i, w in enumerate(det_col_widths, 1):
    wd.column_dimensions[get_column_letter(i)].width = w

wd.freeze_panes = f"A{D_DATA}"
wd.auto_filter.ref = f"A{D_HDR}:{get_column_letter(len(det_headers))}{D_DATA + len(results) - 1}"

wb.save(out_path)
print(f"Saved: {out_path}")
