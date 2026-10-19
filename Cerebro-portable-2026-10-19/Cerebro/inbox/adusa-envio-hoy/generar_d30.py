#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el D30 (viernes 16 oct 2026) a partir del D29 del 13 oct.

Semanal. Periodo que ve Jim: del 9 oct al 16 oct.

Verificado en Jira en vivo el 16 oct:
  Cierres (2): RCS-29026 el 10/13, RCS-29399 el 10/14
  Nuevos: ninguno. Barrido de RCS creados 10/13 a 10/16, cero con reporter adusa.com
  Vuelven dos a Pending Ahold verification: RCS-28693 y RCS-29234, las dos con fix
  liberado y esperando que Ahold lo verifique
  Pasan a Awaiting Ahold input: RCS-29315 y RCS-29164
  No mueven fecha: RCS-28683 y RCS-29000, los dos son espejo de Automation for Jira
  del ticket de ingenieria

  Contexto del periodo: ADUSA tiene un incidente de seguridad desde el fin de semana
  del 10/10 que les limita el acceso por SSO a aplicaciones externas. Jim lo escribio
  en RCS-29234 y en RCS-29164, Rodney en RCS-27077 y Becky en RCS-29300.
"""

import sys, os, re, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    AQUI, "Appriss ADUSA-Support-Tracker - 10132026 D29.xlsx")
DESTINO = os.path.join(AQUI, "Appriss ADUSA-Support-Tracker - 10162026 D30.xlsx")
if not os.path.exists(ORIGEN):
    sys.exit("No encuentro el D29 en:\n  %s" % ORIGEN)

FILL = {"Resolved": "FF7DE2DD", "Active with Appriss": "FFFBC521",
        "Awaiting Ahold input": "FF7C6BE8", "Pending Ahold verification": "FF528DF5"}
def D(d): return datetime.datetime(2026, 10, d)

wb = openpyxl.load_workbook(ORIGEN)
ws = wb["Support tracker"]


def need(clave):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == clave:
            return r
    sys.exit("No encontre %s. Revisa que el archivo sea el D29." % clave)


def categoria(r, cat):
    c = ws.cell(row=r, column=4); c.value = cat
    v = c.fill
    c.fill = PatternFill(fill_type=v.fill_type, fgColor=FILL[cat], bgColor=v.bgColor)


def anexa(r, linea):
    c = ws.cell(row=r, column=10)
    c.value = (c.value or "") + "\n" + linea


def alto(r, h):
    previo = ws.row_dimensions[r].height or 0
    ws.row_dimensions[r].height = max(previo, h)


ws["A2"] = ("Status as of Friday, October 16, 2026. "
            "Ticket states verified directly against Jira.")

# ===================== CIERRES =====================

r = need("RCS-29026")
ws.cell(row=r, column=3).value = "Completed"
categoria(r, "Resolved")
ws.cell(row=r, column=7).value = D(13)
ws.cell(row=r, column=9).value = (
    "The alert for that case did not reach its recipients because one person appeared twice on the "
    "distribution, once on the alert's own email list and once through their store role, with the "
    "two addresses written differently. Appriss reviewed every alert sent this year for the same "
    "pattern and shared the other cases where it appears. Both sides agreed that keeping each "
    "person in one place only is what avoids it. Leigh Ann confirmed on 10/13 that the ticket "
    "could be closed, and it closed the same day.")
anexa(r, "10/13/2026 - We are good to close this one.")

r = need("RCS-29399")
ws.cell(row=r, column=3).value = "Closed"
categoria(r, "Resolved")
ws.cell(row=r, column=7).value = D(14)
ws.cell(row=r, column=9).value = (
    "Appriss checked both instances and shared hour by hour counts for 10/11 through 10/13. On the "
    "Delhaize side the files arrived normally and the load was delayed by an error in the step "
    "that processes the recall extract files, which was restarted, and a correction has been "
    "prepared so a similar file does not stop the load again. On the Ahold side every file that "
    "arrived was processed without errors, and what changed was the content: far fewer rows per "
    "file from 10/12 until the afternoon of 10/13, when volumes returned to normal. Ahold "
    "confirmed on 10/14 that it had regained access to the application and that this answered the "
    "ticket, and it closed the same day.")
anexa(r, "10/13/2026 - Our data processing server shows Delhaize TLOG data is flowing and our "
         "SFTP processes are running normally. On the Ahold side we have a known collection issue "
         "being addressed.")
anexa(r, "10/14/2026 - We now have alternate access to the application and can see this "
         "ourselves again. This resolves the issue in this ticket.")
alto(r, 300)

# ===================== CAMBIOS DE COLUMNA =====================

r = need("RCS-28693")
ws.cell(row=r, column=3).value = "Under Customer Review"
categoria(r, "Pending Ahold verification")
ws.cell(row=r, column=5).value = "Vlad Bogus"
ws.cell(row=r, column=7).value = D(13)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Appriss released a correction for how user hierarchies are read, and asked Ahold on 10/13 to "
    "have Harrison Xu sign in again and report the result. Ahold replied the same day that it "
    "cannot run that check while access is restricted on its side, and asked to keep the ticket "
    "open until it can.")
anexa(r, "10/13/2026 - Appriss released a fix and asked Ahold to have Harrison Xu try signing in "
         "again.")
anexa(r, "10/13/2026 - I am unable to fulfill this request at the moment. Please keep this open "
         "and I will address it when I am able to.")
alto(r, 409)

r = need("RCS-29234")
ws.cell(row=r, column=3).value = "Under Customer Review"
categoria(r, "Pending Ahold verification")
ws.cell(row=r, column=5).value = "Sebastian Garces"
ws.cell(row=r, column=7).value = D(16)
ws.cell(row=r, column=9).value = (
    "The application was restored on 10/2 and no further reports have come in on the ticket since. "
    "The correction went to production on 10/16 as part of this week's scheduled release. Appriss "
    "asked Ahold the same day to confirm that cases save as expected. Ahold replied that it cannot "
    "test while access to the platform is restricted on its side, and expects to come back once "
    "access is restored.")
anexa(r, "10/16/2026 - Appriss confirmed the correction is in production and asked Ahold to "
         "verify.")
anexa(r, "10/16/2026 - We are unable to access LPMS at the moment, so we could not test this. We "
         "hope to have access restored at the beginning of the week and will provide feedback.")
alto(r, 230)

r = need("RCS-29315")
ws.cell(row=r, column=3).value = "Under Customer Review"
categoria(r, "Awaiting Ahold input")
ws.cell(row=r, column=7).value = D(16)
ws.cell(row=r, column=9).value = (
    "Appriss compared this user's sign in against a working one and found that three details the "
    "platform expects are missing from hers: employee ID, title and global ID. The rest arrive "
    "correctly. Appriss asked Ahold on 10/14 to have its IT team complete those attributes on her "
    "profile, then have her sign out fully and sign in again, and report the result. Appriss "
    "followed up on 10/16 and is waiting on that.")
anexa(r, "10/14/2026 - Appriss found three details missing from her sign in, employee ID, title "
         "and global ID, and asked Ahold IT to complete them and retest.")
alto(r, 240)

r = need("RCS-29164")
ws.cell(row=r, column=3).value = "Work in progress"
categoria(r, "Awaiting Ahold input")
ws.cell(row=r, column=7).value = D(15)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Appriss followed up on 10/15. Ahold replied the same day that its LPMS team will review and "
    "respond once access to the platform is restored on its side.")
anexa(r, "10/15/2026 - Our LPMS Support team will review and respond once access is restored.")
alto(r, 210)

# ===================== ACTIVIDAD SIN CAMBIO DE COLUMNA =====================

r = need("RCS-29376")
ws.cell(row=r, column=3).value = "Work in progress"
ws.cell(row=r, column=5).value = "Juan Ramirez"
ws.cell(row=r, column=7).value = D(13)
ws.cell(row=r, column=9).value = (
    "Appriss reviewed the screen against the Ahold tenant on 10/13 and found that the button only "
    "appears on accounts that have never signed in through single sign on. Once an account signs "
    "in that way it is treated as an SSO account and the button is not shown. On the reporter's "
    "own records the two accounts she signs in with do not show it, and the two older accounts "
    "that have never been used do. Appriss is preparing the full answer, including where else the "
    "option can appear.")
anexa(r, "10/13/2026 - Appriss confirmed the option only appears on accounts that have never "
         "signed in through single sign on.")
alto(r, 225)

r = need("RCS-28667")
ws.cell(row=r, column=3).value = "Work in progress"
ws.cell(row=r, column=5).value = "Megha Rathi"
ws.cell(row=r, column=7).value = D(16)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " The change went to production on 10/16 as part of this week's scheduled release. Appriss "
    "will come back to agree with Ahold how it is checked on the Ahold tenant.")
anexa(r, "10/16/2026 - The change was released to production.")
alto(r, 409)

r = need("RCS-29108")
ws.cell(row=r, column=7).value = D(16)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Appriss followed up on 10/16 and is waiting on that confirmation.")
anexa(r, "10/16/2026 - Appriss followed up on the confirmation requested on 10/12.")
alto(r, 220)

r = need("RCS-29300")
ws.cell(row=r, column=3).value = "Work in progress"
categoria(r, "Active with Appriss")
ws.cell(row=r, column=5).value = "Mauricio Zapata"
ws.cell(row=r, column=7).value = D(15)
ws.cell(row=r, column=9).value = (
    "Appriss checked the sample and then the source file for that trading day. The receipt follows "
    "the item sequence number and the audit tab follows scan time, and the two differ because the "
    "sequence numbers in the file arrive out of scan order. The same pattern appears on 392 "
    "tickets across 54 stores in that one file. Ahold explained on 10/14 that its users rely on "
    "scan order to review video alongside the transaction, so the order on the receipt matters for "
    "the investigation itself. Appriss is reviewing how the receipt order is built and whether it "
    "can follow scan time, and is preparing a recent example for the Ahold POS team as requested "
    "on 10/15.")
anexa(r, "10/14/2026 - Our users are accustomed to scan order and rely on it alongside video, so "
         "a different order on the receipt can lead to false positives when investigating.")
anexa(r, "10/15/2026 - Please provide a recent example we can pass to the POS team, as they do "
         "not keep data history for long.")
alto(r, 290)

r = need("RCS-29355")
ws.cell(row=r, column=3).value = "Work in progress"
ws.cell(row=r, column=7).value = D(16)
ws.cell(row=r, column=9).value = (
    "Appriss reviewed the files before loading anything and confirmed on 10/16 that the delivery "
    "is correct on the Ahold side: the store and the tickets in question are present in the test "
    "deliveries, and the recall extract matches the file attached to the ticket. What is missing "
    "is that those test deliveries are not being loaded into the test environment, and Appriss is "
    "reviewing why. Appriss cannot commit to a date yet. On the separate point of sale file, there "
    "is no feed set up for it, and Appriss will confirm the next steps once Ahold has the mapping, "
    "the frequency and the file naming from its project team.")
anexa(r, "10/14/2026 - The target is the QA and test environment. That data has been arriving four "
         "times a day since 9/18. The separate point of sale file comes from the project team and "
         "its mapping is still being worked through on our side.")
anexa(r, "10/16/2026 - Appriss confirmed the test deliveries are correct and that they are not "
         "being loaded into the test environment, and is reviewing why.")
alto(r, 265)

r = need("RCS-27077")
ws.cell(row=r, column=7).value = D(15)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Ahold replied on 10/15 that it will validate the data once access to Secure is restored on "
    "its side.")
anexa(r, "10/15/2026 - We hope to have our access restored shortly, which will let us resume "
         "access to Secure to validate the data.")
alto(r, 300)

r = need("RCS-28281")
ws.cell(row=r, column=7).value = D(13)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Ahold reported on 10/13 that it has no process today for recovering data older than ten days "
    "and is working with its vendors on one. Appriss is keeping the ticket open while that work "
    "runs, and will verify the incoming files and confirm the loaded counts once the August data "
    "is resent.")
anexa(r, "10/13/2026 - We do not currently have a process for recovering data past 10 days. I am "
         "looking into a solution, but it will take some time to develop.")
anexa(r, "10/13/2026 - Appriss confirmed the ticket stays open until the resend lands.")
alto(r, 409)

# ===================== RESUMEN =====================

def fres(nombre):
    for r in range(7, 18):
        if ws.cell(row=r, column=1).value == nombre:
            return r
    sys.exit("No encontre '%s' en el resumen." % nombre)

#                                           D   E   F   G
objetivo = {
 "Login and access":                        (4, 1, 1, 1),
 "User accounts and duplicates":            (6, 0, 0, 0),
 "Store and brand restrictions":            (8, 0, 0, 1),
 "Case entry and assignment":              (11, 1, 3, 2),
 "Attachments and file handling":           (5, 0, 0, 0),
 "Reporting and exports":                   (7, 0, 0, 2),
 "Performance and timeouts":                (5, 0, 0, 0),
 "Platform configuration and notifications":(5, 0, 0, 0),
 "Store file and location data":            (4, 0, 0, 0),
 "Transaction, tender and POS data":        (5, 0, 2, 3),
 "Change requests and scoping":             (1, 0, 0, 0),
}
for tema, (d, e, f, g) in objetivo.items():
    r = fres(tema)
    ws.cell(row=r, column=4).value = d
    ws.cell(row=r, column=5).value = e
    ws.cell(row=r, column=6).value = f
    ws.cell(row=r, column=7).value = g

wb.save(DESTINO)

# ===================== VERIFICACION =====================
ws0 = openpyxl.load_workbook(ORIGEN)["Support tracker"]
ws2 = openpyxl.load_workbook(DESTINO)["Support tracker"]
print("\nGuardado:", os.path.basename(DESTINO)); print(ws2["A2"].value, "\n")
tot = [0, 0, 0, 0]
for r in range(7, 18):
    for i, c in enumerate((4, 5, 6, 7)):
        tot[i] += ws2.cell(row=r, column=c).value or 0
print("resueltos %d | pend verif %d | esperando Ahold %d | activos %d | total %d"
      % (tot[0], tot[1], tot[2], tot[3], sum(tot)))

esp = {
 "RCS-28693": ("Under Customer Review",        "Pending Ahold verification"),
 "RCS-29315": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-29376": ("Work in progress",             "Active with Appriss"),
 "RCS-28667": ("Work in progress",             "Active with Appriss"),
 "RCS-28683": ("Accepted into Development",    "Active with Appriss"),
 "RCS-28849": ("Accepted into Development",    "Active with Appriss"),
 "RCS-29108": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-29131": ("Work in progress",             "Awaiting Ahold input"),
 "RCS-29164": ("Work in progress",             "Awaiting Ahold input"),
 "RCS-29234": ("Under Customer Review",        "Pending Ahold verification"),
 "RCS-28824": ("Work in progress",             "Active with Appriss"),
 "RCS-29000": ("Accepted into Development",    "Active with Appriss"),
 "RCS-28281": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-27077": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-29063": ("Waiting for Software Release", "Active with Appriss"),
 "RCS-29300": ("Work in progress",             "Active with Appriss"),
 "RCS-29355": ("Work in progress",             "Active with Appriss"),
}
cerrados_hoy = ["RCS-29026", "RCS-29399"]
ok, vistos, claves = True, set(), []
print("Abiertos:")
for r in range(21, ws2.max_row + 1):
    k = ws2.cell(row=r, column=1).value
    if not (k and str(k).startswith("RCS-")):
        continue
    claves.append(k)
    cat = ws2.cell(row=r, column=4).value
    if cat == "Resolved":
        continue
    est = ws2.cell(row=r, column=3).value
    vistos.add(k)
    mal = k in esp and (est, cat) != esp[k]
    ok = ok and not mal
    print("  %s fila %-4d %-11s %-28s %-27s %s" % (
        "!!" if mal else "  ", r, k, est, cat, ws2.cell(row=r, column=5).value))
if set(esp) - vistos:
    ok = False; print("\nFaltan abiertos:", ", ".join(sorted(set(esp) - vistos)))
if vistos - set(esp):
    ok = False; print("\nAbiertos inesperados:", ", ".join(sorted(vistos - set(esp))))
print("\nCerrados en el periodo:")
for k in cerrados_hoy:
    for r in range(21, ws2.max_row + 1):
        if str(ws2.cell(row=r, column=1).value).strip() == k:
            c4 = ws2.cell(row=r, column=4)
            bien = c4.value == "Resolved" and c4.fill.fgColor.rgb == FILL["Resolved"]
            ok = ok and bien
            print("  %s fila %-4d %-11s %-11s %-9s fill=%s" % (
                "  " if bien else "!!", r, k, ws2.cell(row=r, column=3).value,
                c4.value, c4.fill.fgColor.rgb))
            break
    else:
        ok = False; print("  !! %s no esta en el documento" % k)
if sum(tot) != 78 or tot[0] != 61:
    ok = False; print("\nTotales mal, deberian ser 78 y 61 resueltos.")
d = sorted({k for k in claves if claves.count(k) > 1})
if d:
    ok = False; print("\nDuplicados:", ", ".join(d))
print("\ntickets en el documento: %d" % len(claves))

def mapa(w):
    m = {}
    for r in range(21, w.max_row + 1):
        v = w.cell(row=r, column=1).value
        if v and str(v).startswith("RCS-"):
            m[str(v)] = r
    return m
m0, m2 = mapa(ws0), mapa(ws2)
print("\nDiff contra el D29:")
cambios, tocados = 0, set()
for k, r0 in m0.items():
    r2 = m2.get(k)
    if r2 is None:
        ok = False; print("  !! %s desaparecio" % k); continue
    for c in range(1, 11):
        if ws0.cell(row=r0, column=c).value != ws2.cell(row=r2, column=c).value:
            cambios += 1; tocados.add(k)
print("  tickets tocados: %s" % ", ".join(sorted(tocados)))
print("  filas nuevas: %s" % (", ".join(sorted(set(m2) - set(m0))) or "ninguna"))
print("  celdas cambiadas: %d" % cambios)
malas = 0
for k, r2 in m2.items():
    if k in m0 and (ws2.row_dimensions[r2].height or 0) < (ws0.row_dimensions[m0[k]].height or 0):
        malas += 1; print("  !! %s bajo de alto" % k)
print("  alturas mal desplazadas: %d" % malas)
ok = ok and malas == 0
print("\n" + ("Todo cuadra." if ok else "Revisa lo marcado."))
