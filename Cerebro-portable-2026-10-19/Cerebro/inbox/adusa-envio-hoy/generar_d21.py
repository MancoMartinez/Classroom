#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el D21 (lunes 14 sep 2026) a partir del D20.

Cambios contra la version que llego por chat:
  1. nueva() ahora desplaza las alturas de fila. openpyxl no lo hace en
     insert_rows, asi que sin esto cada fila debajo del punto de insercion
     se quedaba con la altura de la fila que estaba antes en ese numero.
  2. RCS-28911 recoge el segundo comentario de April Valerio del 9/14 19:39.
"""

import sys, os, re, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    AQUI, "Appriss ADUSA-Support-Tracker - 09112026 D20.xlsx")
DESTINO = os.path.join(AQUI, "Appriss ADUSA-Support-Tracker - 09142026 D21.xlsx")
if not os.path.exists(ORIGEN):
    sys.exit("No encuentro el D20 en:\n  %s" % ORIGEN)

FILL = {"Resolved": "FF7DE2DD", "Active with Appriss": "FFFBC521",
        "Awaiting Ahold input": "FF7C6BE8", "Pending Ahold verification": "FF528DF5"}
BLANCO, BANDA, HOY = "FFFFFFFF", "FFF4F6F9", datetime.datetime(2026, 9, 14)

wb = openpyxl.load_workbook(ORIGEN)
ws = wb["Support tracker"]


def need(clave):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == clave:
            return r
    sys.exit("No encontre %s. Revisa que el archivo sea el D20." % clave)


def categoria(r, cat):
    c = ws.cell(row=r, column=4); c.value = cat
    v = c.fill
    c.fill = PatternFill(fill_type=v.fill_type, fgColor=FILL[cat], bgColor=v.bgColor)


def anexa(r, linea):
    c = ws.cell(row=r, column=10)
    c.value = (c.value or "") + "\n" + linea


def bloque(nombre):
    hdr = None
    for r in range(23, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == nombre:
            hdr = r; break
    if hdr is None:
        sys.exit("No encontre el bloque '%s'." % nombre)
    ultima = hdr
    for r in range(hdr + 1, ws.max_row + 2):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().startswith("RCS-"):
            ultima = r
        elif v:
            break
    return ultima


def nueva(nombre_tema, vals, alto):
    plantilla = bloque(nombre_tema)
    dest = plantilla + 1

    guardadas = [str(m) for m in list(ws.merged_cells.ranges) if m.min_row >= dest]
    for s in guardadas:
        ws.unmerge_cells(s)

    # openpyxl no desplaza row_dimensions en insert_rows, hay que hacerlo a mano
    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if d.height is not None and r >= dest}

    ws.insert_rows(dest, amount=1)

    for s in guardadas:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", s)
        ws.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1,
                                      m.group(3), int(m.group(4)) + 1))

    for r in sorted(alturas, reverse=True):
        ws.row_dimensions[r + 1].height = alturas[r]
    ws.row_dimensions[dest].height = alto

    for c in range(1, 11):
        ws.cell(row=dest, column=c)._style = copy(ws.cell(row=plantilla, column=c)._style)
    fg = ws.cell(row=plantilla, column=1).fill.fgColor
    previa = fg.rgb if fg.type == "rgb" else BLANCO
    banda = BLANCO if previa == BANDA else BANDA
    for c in range(1, 11):
        ws.cell(row=dest, column=c).fill = PatternFill(
            fill_type="solid", fgColor=FILL[vals[3]] if c == 4 else banda)
    for i, v in enumerate(vals, start=1):
        ws.cell(row=dest, column=i).value = v
    for c in (6, 7):
        ws.cell(row=dest, column=c).number_format = \
            ws.cell(row=plantilla, column=c).number_format


ws["A2"] = ("Status as of Monday, September 14, 2026. "
            "Ticket states verified directly against Jira.")

# ---------- RCS-28752, cerrado hoy ----------
r = need("RCS-28752")
ws.cell(row=r, column=3).value = "Completed"
categoria(r, "Resolved")
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "The process that assigns a location to a new user matched on store number without the "
    "division, so at a store number shared by Food Lion and Stop and Shop it returned the wrong "
    "one. Appriss corrected the view and applied the fix to the Ahold environment. The location "
    "assignment on the account was correct all along, so nothing changed on the user. Leigh Ann "
    "put the user on it and came back on 9/14, and Appriss closed the ticket the same afternoon. "
    "The wider code change continues separately so the correction reaches every client.")
anexa(r, "9/14/2026 - Ahold came back on the retest and the ticket closed as resolved.")

# ---------- RCS-28849, pasa a change request ----------
r = need("RCS-28849")
ws.cell(row=r, column=3).value = "Accepted into Development"
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "Appriss compared the assign and the reject paths. The reject path sends the user to the "
    "confirmation page, the assign path does not, and it has not carried that screen for around "
    "three years. RCS-28123 reached the same conclusion. Appriss reviewed it with engineering on "
    "9/14 and accepted it as a change request rather than a defect, so adding the screen is a "
    "product enhancement. Appriss will update this ticket as that work progresses.")
anexa(r, "9/14/2026 - Accepted as a change request. The confirmation screen on the assign path is "
         "current behaviour by design, so adding it goes through product rather than as a fix.")

# ---------- RCS-28281, sigue esperando a Ahold ----------
r = need("RCS-28281")
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "Appriss posted the full count per trading day on 8/27. Ahold reported on 8/31 that a good "
    "deal of recovery had occurred, with two gaps left on its view, the 9 PM to 10 PM hour on "
    "8/12 and a handful of trading days sitting 100,000 to 200,000 transactions below average. "
    "Appriss asked Rodney on 9/7 to have Becky run the same 10 minute check she used on 8/20 "
    "against 8/11, 8/12, 8/19, 8/20, 8/22 and 8/25, so Appriss can confirm which batches still "
    "need resubmission. That reply has not arrived and the ticket has sat seven days on it.")

# ---------- RCS-28910, nuevo ----------
nueva("Case entry and assignment", [
    "RCS-28910",
    "New case type cannot be closed and its workbench view differs",
    "Under Support Review", "Active with Appriss", "Unassigned", HOY, HOY,
    "The Elevator Entrapment case type was created shortly before the move to the new URL and two "
    "things do not behave like the other case types. A case of that type cannot be closed from the "
    "workbench under Update Status, even though the workflow allows it. The workbench for that "
    "type also does not show the side bar or the case information bar that other types show. "
    "Raised 9/14 with screenshots.",
    "Open with Appriss. Raised today and under support review. Appriss is picking up ownership and "
    "will confirm the next step.",
    "Ticket opened 9/14/2026\n9/14/2026 - No ability to close the case from the workbench, and the "
    "workbench is missing the side bar and case info bar that another case type has.",
], 135)

# ---------- RCS-28911, nuevo ----------
nueva("Store file and location data", [
    "RCS-28911",
    "Daily location file not processed",
    "Under Support Review", "Active with Appriss", "Unassigned", HOY, HOY,
    "The daily location update file did not process on the morning of 9/14. Ahold reports the file "
    "left its side and asks Appriss to confirm receipt. Ahold followed up the same evening and "
    "notes the file carries a large batch of store changes still to be applied in LPMS.",
    "Open with Appriss. Raised today and under support review. Appriss is checking whether the "
    "file arrived and will come back to Ahold with what it finds.",
    "Ticket opened 9/14/2026\n9/14/2026 - The daily location update file was not processed this "
    "morning.\n9/14/2026 - Please check to confirm you received the file. It appears to be sent "
    "from our side.\n9/14/2026 - Still looking for an update on processing of the store file. "
    "There were lots of changes that needed to be updated in LPMS.",
], 135)

# ---------- resumen ----------
def fres(nombre):
    for r in range(7, 18):
        if ws.cell(row=r, column=1).value == nombre:
            return r
    sys.exit("No encontre '%s' en el resumen." % nombre)

r = fres("Store and brand restrictions"); ws.cell(row=r, column=4).value = 8; ws.cell(row=r, column=5).value = 0
ws.cell(row=fres("Case entry and assignment"), column=7).value = 3
ws.cell(row=fres("Store file and location data"), column=7).value = 2

wb.save(DESTINO)

# ---------- verificacion ----------
ws2 = openpyxl.load_workbook(DESTINO)["Support tracker"]
print("\nGuardado:", os.path.basename(DESTINO)); print(ws2["A2"].value, "\n")
tot = [0, 0, 0, 0]
for r in range(7, 18):
    for i, c in enumerate((4, 5, 6, 7)):
        tot[i] += ws2.cell(row=r, column=c).value or 0
print("resueltos %d | pend verif %d | esperando Ahold %d | activos %d | total %d"
      % (tot[0], tot[1], tot[2], tot[3], sum(tot)))
esp = {"RCS-28693": ("Under Engineering Review", "Active with Appriss"),
       "RCS-28667": ("Under Engineering Review", "Active with Appriss"),
       "RCS-28683": ("Under Engineering Review", "Active with Appriss"),
       "RCS-28849": ("Accepted into Development", "Active with Appriss"),
       "RCS-28910": ("Under Support Review", "Active with Appriss"),
       "RCS-28824": ("Work in progress", "Active with Appriss"),
       "RCS-28589": ("Accepted into Development", "Active with Appriss"),
       "RCS-28911": ("Under Support Review", "Active with Appriss"),
       "RCS-28281": ("Under Customer Review", "Awaiting Ahold input")}
ok, vistos, claves = True, set(), []
print("\nAbiertos:")
for r in range(23, ws2.max_row + 1):
    k = ws2.cell(row=r, column=1).value
    if not (k and str(k).startswith("RCS-")):
        continue
    claves.append(k)
    cat = ws2.cell(row=r, column=4).value
    if cat == "Resolved":
        continue
    est = ws2.cell(row=r, column=3).value
    vistos.add(k)
    mal = k in esp and (est, cat) != esp[k]
    ok = ok and not mal
    print("  %s fila %-4d %-11s %-26s %s" % ("!!" if mal else "  ", r, k, est, cat))
if set(esp) - vistos:
    ok = False; print("\nFaltan:", ", ".join(sorted(set(esp) - vistos)))
if vistos - set(esp):
    ok = False; print("\nAbiertos inesperados:", ", ".join(sorted(vistos - set(esp))))
if sum(tot) != 54 or tot[0] != 45:
    ok = False; print("\nTotales mal, deberian ser 54 y 45 resueltos.")
d = sorted({k for k in claves if claves.count(k) > 1})
if d:
    ok = False; print("\nDuplicados:", ", ".join(d))
print("\n" + ("Todo cuadra." if ok else "Revisa lo marcado."))
