#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el D29 (martes 13 oct 2026) a partir del D28 del 9 oct.

Tracker fuera de cadencia, pedido por J.R. Teagle para compartir con Carl G.
Periodo cubierto: 9 al 13 oct.

Verificado en Jira en vivo el 13 oct:
  Cierres: ninguno.
  Cambian de estado (5): RCS-29315, RCS-29376, RCS-28667, RCS-29108, RCS-27077
  Cambian de columna (2): RCS-29108 y RCS-27077 pasan a Awaiting Ahold input
  Nuevos (1): RCS-29399
  Correccion del D28: RCS-28667 ya estaba en Accepted into Development desde el 10/2,
  el D28 lo reporto como Under Engineering Review.
  No mueven fecha: RCS-29000 (10/12 es espejo de Automation for Jira del ticket ENG),
  RCS-28281 (10/11 es aviso automatico de 60 dias; la fecha sube a 10/6 por el
  seguimiento real de Juan), RCS-28667 (10/13 es espejo de Automation for Jira).
"""

import sys, os, re, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    AQUI, "Appriss ADUSA-Support-Tracker - 10092026 D28.xlsx")
DESTINO = os.path.join(AQUI, "Appriss ADUSA-Support-Tracker - 10132026 D29.xlsx")
if not os.path.exists(ORIGEN):
    sys.exit("No encuentro el D28 en:\n  %s" % ORIGEN)

FILL = {"Resolved": "FF7DE2DD", "Active with Appriss": "FFFBC521",
        "Awaiting Ahold input": "FF7C6BE8", "Pending Ahold verification": "FF528DF5"}
BLANCO, BANDA = "FFFFFFFF", "FFF4F6F9"
def D(d): return datetime.datetime(2026, 10, d)

wb = openpyxl.load_workbook(ORIGEN)
ws = wb["Support tracker"]


def need(clave):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == clave:
            return r
    sys.exit("No encontre %s. Revisa que el archivo sea el D28." % clave)


def categoria(r, cat):
    c = ws.cell(row=r, column=4); c.value = cat
    v = c.fill
    c.fill = PatternFill(fill_type=v.fill_type, fgColor=FILL[cat], bgColor=v.bgColor)


def anexa(r, linea):
    c = ws.cell(row=r, column=10)
    c.value = (c.value or "") + "\n" + linea


def alto(r, h):
    # una fila existente nunca baja de alto: la columna J solo crece
    previo = ws.row_dimensions[r].height or 0
    ws.row_dimensions[r].height = max(previo, h)


def bloque(nombre):
    hdr = None
    for r in range(21, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == nombre:
            hdr = r; break
    if hdr is None:
        sys.exit("No encontre el bloque '%s'." % nombre)
    ultima = hdr
    for r in range(hdr + 1, ws.max_row + 2):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().startswith("RCS-"):
            ultima = r
        elif v:
            break
    return ultima


def nueva(nombre_tema, vals, h):
    plantilla = bloque(nombre_tema)
    dest = plantilla + 1
    guardadas = [str(m) for m in list(ws.merged_cells.ranges) if m.min_row >= dest]
    for s in guardadas:
        ws.unmerge_cells(s)
    # openpyxl no desplaza row_dimensions en insert_rows, hay que hacerlo a mano
    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if d.height is not None and r >= dest}
    ws.insert_rows(dest, amount=1)
    for s in guardadas:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", s)
        ws.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1,
                                      m.group(3), int(m.group(4)) + 1))
    for r in sorted(alturas, reverse=True):
        ws.row_dimensions[r + 1].height = alturas[r]
    ws.row_dimensions[dest].height = h
    for c in range(1, 11):
        ws.cell(row=dest, column=c)._style = copy(ws.cell(row=plantilla, column=c)._style)
    fg = ws.cell(row=plantilla, column=1).fill.fgColor
    previa = fg.rgb if fg.type == "rgb" else BLANCO
    banda = BLANCO if previa == BANDA else BANDA
    for c in range(1, 11):
        ws.cell(row=dest, column=c).fill = PatternFill(
            fill_type="solid", fgColor=FILL[vals[3]] if c == 4 else banda)
    for i, v in enumerate(vals, start=1):
        ws.cell(row=dest, column=i).value = v
    for c in (6, 7):
        ws.cell(row=dest, column=c).number_format = \
            ws.cell(row=plantilla, column=c).number_format


ws["A2"] = ("Status as of Tuesday, October 13, 2026. "
            "Ticket states verified directly against Jira.")

# ===================== CAMBIOS DE ESTADO =====================

r = need("RCS-29315")
ws.cell(row=r, column=3).value = "Under Engineering Review"
ws.cell(row=r, column=7).value = D(13)
ws.cell(row=r, column=9).value = (
    "Appriss traced the sign in through the single sign on log and confirmed on 10/13 that the "
    "user's sign in reaches the platform, carrying her role and location details. No account is "
    "created for her after that point, so the behaviour sits in how the platform handles the sign "
    "in rather than on the Ahold side. Appriss passed the case to engineering the same day. "
    "Nothing is needed from Ahold while that review runs.")
anexa(r, "10/13/2026 - Appriss confirmed the sign in reaches the platform with her role and "
         "location details, and no account is created. Passed to engineering, no action needed "
         "from Ahold for now.")
alto(r, 190)

r = need("RCS-29376")
ws.cell(row=r, column=3).value = "Under Delivery Review"
ws.cell(row=r, column=5).value = "David Parraga"
ws.cell(row=r, column=7).value = D(12)
ws.cell(row=r, column=9).value = (
    "Appriss reproduced the button on the Ahold tenant and confirmed it appears on the user record "
    "in the Admin users screen. The case moved on 10/12 from engineering to the delivery team for "
    "initial review, since what is needed is a configuration answer rather than a platform "
    "change. Appriss will come back on whether the option can be removed for Ahold and where else "
    "it appears.")
anexa(r, "10/12/2026 - Appriss moved the case to the delivery team for initial review.")
alto(r, 175)

r = need("RCS-28667")
ws.cell(row=r, column=3).value = "Accepted into Development"
ws.cell(row=r, column=7).value = D(2)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Engineering accepted the change into development on 10/2. Appriss will update this ticket "
    "as that work progresses, and will come back to agree how the change is tested before it is "
    "released.")
anexa(r, "10/2/2026 - Engineering accepted the change into development.")
alto(r, 409)

r = need("RCS-29108")
ws.cell(row=r, column=3).value = "Under Customer Review"
categoria(r, "Awaiting Ahold input")
ws.cell(row=r, column=7).value = D(12)
ws.cell(row=r, column=9).value = (
    "Appriss checked the reassignment notice and the repeated emails separately, since each comes "
    "from a different part of the platform. The setting that sends a notice on reassignment is "
    "switched on for Ahold, the user has an email address on her record, and the correction behind "
    "the missing notice is already in production. Appriss asked Ahold on 10/12 to confirm whether "
    "that user is still not receiving the notices, and for the case number and the time of an "
    "assignment if she is, so the remaining half can be traced.")
anexa(r, "10/12/2026 - Appriss confirmed the correction is in production and asked Ahold to check "
         "whether the user is still not receiving the notices.")
alto(r, 205)

r = need("RCS-27077")
ws.cell(row=r, column=3).value = "Under Customer Review"
categoria(r, "Awaiting Ahold input")
ws.cell(row=r, column=7).value = D(12)
ws.cell(row=r, column=9).value = (
    "The Delhaize side of this request is resolved and deployed, and all September and October "
    "test files were loaded by 10/7. Appriss received the Ahold test data on 10/8 and confirmed on "
    "10/12 that all 76 files loaded, with none skipped or rejected, for 12,758 transactions. Ahold "
    "confirmed on 10/9 that this is the last load needed unless the testing turns something up. "
    "Appriss is waiting on Ahold to confirm whether this ticket can be closed, with the Ahold "
    "requirements tracked on RCS-29063 so each banner can move on its own.")
anexa(r, "10/12/2026 - Appriss loaded all 76 Ahold test files, 12,758 transactions, none skipped "
         "or rejected.")
alto(r, 280)

# ===================== ACTIVIDAD SIN CAMBIO DE ESTADO =====================

r = need("RCS-29300")
ws.cell(row=r, column=7).value = D(13)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Appriss followed up on 10/13 and is waiting on that review before taking the case further.")
anexa(r, "10/13/2026 - Appriss followed up on the findings shared on 10/9.")
alto(r, 215)

r = need("RCS-28281")
ws.cell(row=r, column=7).value = D(6)
ws.cell(row=r, column=9).value = (
    ws.cell(row=r, column=9).value +
    " Appriss followed up on 10/6 and is still waiting on the resend of the 8/12 window. The "
    "ticket has been open for more than sixty days and is on the list for review with Appriss "
    "support management.")
anexa(r, "10/6/2026 - Appriss followed up on the 8/12 resend request.")
alto(r, 409)

# ===================== NUEVOS =====================

nueva("Transaction, tender and POS data", [
    "RCS-29399", "Validation requested that transaction data is still arriving",
    "Work in progress", "Active with Appriss", "Luis Hernandez", D(12), D(13),
    "Ahold asked Appriss to confirm whether transaction data is still arriving on both instances "
    "and, if it stopped, when. Ahold reports an event on its side that is limiting access to "
    "externally hosted systems through the normal security route, starting over the weekend of "
    "10/10. No data has been regenerated on the Ahold side. Raised 10/12.",
    "Appriss shared a daily status on 10/12. Ahold reported on 10/13 that some of the Ahold data "
    "is arriving and none of the Delhaize data, and asked for hour by hour counts for 10/11 "
    "through 10/13. Appriss is preparing those counts. Ahold also asked whether Secure can be "
    "reached without single sign on while the normal route is unavailable, and Appriss is "
    "reviewing that request.",
    "Ticket opened 10/12/2026\n"
    "10/12/2026 - We are not currently able to access externally hosted systems through our normal "
    "security protocols. Nothing has been regenerated yet. We are getting a baseline on whether "
    "data is flowing successfully to your servers, and if it stopped, when.\n"
    "10/13/2026 - Some of the Ahold data is coming in, none of the Delhaize data. We would like an "
    "hour by hour for 10/11 through 10/13 to see when the issues started.",
], 215)

# ===================== RESUMEN =====================

def fres(nombre):
    for r in range(7, 18):
        if ws.cell(row=r, column=1).value == nombre:
            return r
    sys.exit("No encontre '%s' en el resumen." % nombre)

#                                           D   E   F   G
objetivo = {
 "Login and access":                        (4, 0, 0, 3),
 "User accounts and duplicates":            (6, 0, 0, 0),
 "Store and brand restrictions":            (8, 0, 0, 1),
 "Case entry and assignment":              (11, 0, 2, 4),
 "Attachments and file handling":           (5, 0, 0, 0),
 "Reporting and exports":                   (7, 0, 0, 2),
 "Performance and timeouts":                (5, 0, 0, 0),
 "Platform configuration and notifications":(4, 0, 0, 1),
 "Store file and location data":            (4, 0, 0, 0),
 "Transaction, tender and POS data":        (4, 0, 3, 3),
 "Change requests and scoping":             (1, 0, 0, 0),
}
for tema, (d, e, f, g) in objetivo.items():
    r = fres(tema)
    ws.cell(row=r, column=4).value = d
    ws.cell(row=r, column=5).value = e
    ws.cell(row=r, column=6).value = f
    ws.cell(row=r, column=7).value = g

wb.save(DESTINO)

# ===================== VERIFICACION =====================
ws0 = openpyxl.load_workbook(ORIGEN)["Support tracker"]
ws2 = openpyxl.load_workbook(DESTINO)["Support tracker"]
print("\nGuardado:", os.path.basename(DESTINO)); print(ws2["A2"].value, "\n")
tot = [0, 0, 0, 0]
for r in range(7, 18):
    for i, c in enumerate((4, 5, 6, 7)):
        tot[i] += ws2.cell(row=r, column=c).value or 0
print("resueltos %d | pend verif %d | esperando Ahold %d | activos %d | total %d"
      % (tot[0], tot[1], tot[2], tot[3], sum(tot)))

esp = {
 "RCS-28693": ("Under Engineering Review",     "Active with Appriss"),
 "RCS-29315": ("Under Engineering Review",     "Active with Appriss"),
 "RCS-29376": ("Under Delivery Review",        "Active with Appriss"),
 "RCS-28667": ("Accepted into Development",    "Active with Appriss"),
 "RCS-28683": ("Accepted into Development",    "Active with Appriss"),
 "RCS-28849": ("Accepted into Development",    "Active with Appriss"),
 "RCS-29108": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-29131": ("Work in progress",             "Awaiting Ahold input"),
 "RCS-29164": ("Work in progress",             "Active with Appriss"),
 "RCS-29234": ("Accepted into Development",    "Active with Appriss"),
 "RCS-28824": ("Work in progress",             "Active with Appriss"),
 "RCS-29000": ("Accepted into Development",    "Active with Appriss"),
 "RCS-29026": ("Work in progress",             "Active with Appriss"),
 "RCS-28281": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-27077": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-29063": ("Waiting for Software Release", "Active with Appriss"),
 "RCS-29300": ("Under Customer Review",        "Awaiting Ahold input"),
 "RCS-29355": ("Under Support Review",         "Active with Appriss"),
 "RCS-29399": ("Work in progress",             "Active with Appriss"),
}
ok, vistos, claves = True, set(), []
print("Abiertos:")
for r in range(21, ws2.max_row + 1):
    k = ws2.cell(row=r, column=1).value
    if not (k and str(k).startswith("RCS-")):
        continue
    claves.append(k)
    cat = ws2.cell(row=r, column=4).value
    if cat == "Resolved":
        continue
    est = ws2.cell(row=r, column=3).value
    vistos.add(k)
    mal = k in esp and (est, cat) != esp[k]
    ok = ok and not mal
    print("  %s fila %-4d %-11s %-28s %-27s %s" % (
        "!!" if mal else "  ", r, k, est, cat, ws2.cell(row=r, column=5).value))
if set(esp) - vistos:
    ok = False; print("\nFaltan abiertos:", ", ".join(sorted(set(esp) - vistos)))
if vistos - set(esp):
    ok = False; print("\nAbiertos inesperados:", ", ".join(sorted(vistos - set(esp))))
if sum(tot) != 78 or tot[0] != 59:
    ok = False; print("\nTotales mal, deberian ser 78 y 59 resueltos.")
d = sorted({k for k in claves if claves.count(k) > 1})
if d:
    ok = False; print("\nDuplicados:", ", ".join(d))
print("\ntickets en el documento: %d" % len(claves))

# diff celda por celda contra el D28, con mapeo de filas por clave
def mapa(w):
    m = {}
    for r in range(21, w.max_row + 1):
        v = w.cell(row=r, column=1).value
        if v and str(v).startswith("RCS-"):
            m[str(v)] = r
    return m
m0, m2 = mapa(ws0), mapa(ws2)
print("\nDiff contra el D28:")
cambios = 0
for k, r0 in m0.items():
    r2 = m2.get(k)
    if r2 is None:
        ok = False; print("  !! %s desaparecio" % k); continue
    for c in range(1, 11):
        v0, v2 = ws0.cell(row=r0, column=c).value, ws2.cell(row=r2, column=c).value
        if v0 != v2:
            cambios += 1
            print("  %s col %d" % (k, c))
    h0 = ws0.row_dimensions[r0].height or 0
    h2 = ws2.row_dimensions[r2].height or 0
    if h2 < h0:
        ok = False; print("  !! %s bajo de alto %s -> %s" % (k, h0, h2))
nuevas = set(m2) - set(m0)
print("  filas nuevas:", ", ".join(sorted(nuevas)) or "ninguna")
print("  celdas cambiadas en filas existentes: %d" % cambios)

# alturas desplazadas: toda fila con clave debe conservar su alto
malas = 0
for k, r2 in m2.items():
    if k in m0 and (ws2.row_dimensions[r2].height or 0) < (ws0.row_dimensions[m0[k]].height or 0):
        malas += 1
print("  alturas mal desplazadas: %d" % malas)
ok = ok and malas == 0

print("\n" + ("Todo cuadra." if ok else "Revisa lo marcado."))
