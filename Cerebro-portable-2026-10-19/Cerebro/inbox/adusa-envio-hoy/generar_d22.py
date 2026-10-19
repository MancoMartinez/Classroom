#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera el D22 (martes 15 sep 2026) a partir del D21.

Cambios de hoy, todos verificados en Jira en vivo:
  RCS-28910  a Work in progress, asignado a Vedant Shirude
  RCS-28824  a Under Engineering Review, asignado a Gabor Remenyi
  RCS-28911  a Under Customer Review, asignado a Juan Martinez, pasa a Awaiting Ahold input
  RCS-28944  nuevo, bloque Reporting and exports
  RCS-28281  sin actividad real, no se toca

nueva() desplaza alturas de fila y merges a mano. openpyxl no lo hace en insert_rows.
"""

import sys, os, re, datetime
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    AQUI, "Appriss ADUSA-Support-Tracker - 09142026 D21.xlsx")
DESTINO = os.path.join(AQUI, "Appriss ADUSA-Support-Tracker - 09152026 D22.xlsx")
if not os.path.exists(ORIGEN):
    sys.exit("No encuentro el D21 en:\n  %s" % ORIGEN)

FILL = {"Resolved": "FF7DE2DD", "Active with Appriss": "FFFBC521",
        "Awaiting Ahold input": "FF7C6BE8", "Pending Ahold verification": "FF528DF5"}
BLANCO, BANDA, HOY = "FFFFFFFF", "FFF4F6F9", datetime.datetime(2026, 9, 15)

wb = openpyxl.load_workbook(ORIGEN)
ws = wb["Support tracker"]


def need(clave):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == clave:
            return r
    sys.exit("No encontre %s. Revisa que el archivo sea el D21." % clave)


def categoria(r, cat):
    c = ws.cell(row=r, column=4); c.value = cat
    v = c.fill
    c.fill = PatternFill(fill_type=v.fill_type, fgColor=FILL[cat], bgColor=v.bgColor)


def anexa(r, linea):
    c = ws.cell(row=r, column=10)
    c.value = (c.value or "") + "\n" + linea


def bloque(nombre):
    hdr = None
    for r in range(23, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip() == nombre:
            hdr = r; break
    if hdr is None:
        sys.exit("No encontre el bloque '%s'." % nombre)
    ultima = hdr
    for r in range(hdr + 1, ws.max_row + 2):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().startswith("RCS-"):
            ultima = r
        elif v:
            break
    return ultima


def nueva(nombre_tema, vals, alto):
    plantilla = bloque(nombre_tema)
    dest = plantilla + 1

    guardadas = [str(m) for m in list(ws.merged_cells.ranges) if m.min_row >= dest]
    for s in guardadas:
        ws.unmerge_cells(s)

    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if d.height is not None and r >= dest}

    ws.insert_rows(dest, amount=1)

    for s in guardadas:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", s)
        ws.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1,
                                      m.group(3), int(m.group(4)) + 1))

    for r in sorted(alturas, reverse=True):
        ws.row_dimensions[r + 1].height = alturas[r]
    ws.row_dimensions[dest].height = alto

    for c in range(1, 11):
        ws.cell(row=dest, column=c)._style = copy(ws.cell(row=plantilla, column=c)._style)
    fg = ws.cell(row=plantilla, column=1).fill.fgColor
    previa = fg.rgb if fg.type == "rgb" else BLANCO
    banda = BLANCO if previa == BANDA else BANDA
    for c in range(1, 11):
        ws.cell(row=dest, column=c).fill = PatternFill(
            fill_type="solid", fgColor=FILL[vals[3]] if c == 4 else banda)
    for i, v in enumerate(vals, start=1):
        ws.cell(row=dest, column=i).value = v
    for c in (6, 7):
        ws.cell(row=dest, column=c).number_format = \
            ws.cell(row=plantilla, column=c).number_format


ws["A2"] = ("Status as of Tuesday, September 15, 2026. "
            "Ticket states verified directly against Jira.")

# ---------- RCS-28910, Appriss toma el ticket ----------
r = need("RCS-28910")
ws.cell(row=r, column=3).value = "Work in progress"
ws.cell(row=r, column=5).value = "Vedant Shirude"
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "Appriss took ownership on 9/15 and is working through both points, the case of that type "
    "that cannot be closed from the workbench and the missing side bar and case information bar. "
    "Appriss will come back to Ahold with what it finds.")
anexa(r, "9/15/2026 - Appriss took ownership and started the investigation.")

# ---------- RCS-28824, avanza y queda con ingenieria ----------
r = need("RCS-28824")
ws.cell(row=r, column=3).value = "Under Engineering Review"
ws.cell(row=r, column=5).value = "Gabor Remenyi"
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "Appriss cleared what was blocking the store and cashier profiles on the Delhaize instance. "
    "Only the weekly figures had been rebuilt at first, so Appriss extended that work on 9/15 and "
    "Becky confirmed the Month, Quarter and Year tabs now return charts. Two points stay open. The "
    "cashier profile page still returns no data, so the correction reached the store profile only. "
    "The area map section is still blank on every tab, and Ahold asks whether it can be built from "
    "the location data it already sends, which carries street addresses. The ticket sits with "
    "engineering on both.")
anexa(r, "9/15/2026 - Month, Quarter and Year tabs now show charts.")
anexa(r, "9/15/2026 - The cashier profile page still returns no data, and the area map is still "
         "blank on the profile pages.")

# ---------- RCS-28911, corregido tras el aporte de Ryan Barlow en Teams ----------
r = need("RCS-28911")
ws.cell(row=r, column=3).value = "Under Customer Review"
categoria(r, "Awaiting Ahold input")
ws.cell(row=r, column=5).value = "Juan Martinez"
ws.cell(row=r, column=7).value = HOY
ws.cell(row=r, column=9).value = (
    "The store and location update was not part of the load that ran on the morning of 9/14. The "
    "current read is that the file Ahold sent arrived after that run and was picked up on the next "
    "scheduled one, so the store changes may already be in place. Appriss asked Ahold on 9/15 to "
    "confirm whether those changes are showing in LPMS now, and to hold off on sending another "
    "copy while that is checked, because the files arrive under the same name into the same folder "
    "and a second copy can land on top of one that is still waiting to be picked up. Appriss will "
    "pick this back up if the changes are not there.")
anexa(r, "9/15/2026 - The store and location update was not part of the 9/14 run. It looks like the "
         "file arrived after that run and was picked up on the next one. Could you confirm whether "
         "the store changes are showing in LPMS now, and hold off on sending another copy while "
         "that is checked?")

# ---------- RCS-28944, nuevo ----------
nueva("Reporting and exports", [
    "RCS-28944",
    "Field mappings lost on a rebuilt location listing report",
    "Under Support Review", "Active with Appriss", "Unassigned", HOY, HOY,
    "A location listing report was rebuilt and the mappings that turn binary values into readable "
    "text no longer work. The example given is the Inactive field, which should read Active when "
    "the value is zero and Inactive otherwise. Ahold has tried to recreate the mappings without "
    "success and asks for help setting them up. Raised 9/15.",
    "Open with Appriss. Raised on 9/15 and under support review. Appriss is picking up ownership "
    "and will come back with the steps to rebuild the mappings.",
    "Ticket opened 9/15/2026\n9/15/2026 - Rebuilt a location listing report and the mappings from "
    "binary results to text need to be recreated. One example is Inactive, which should read "
    "Active if the value is 0 and Inactive otherwise.",
], 135)

# ---------- resumen ----------
def fres(nombre):
    for r in range(7, 18):
        if ws.cell(row=r, column=1).value == nombre:
            return r
    sys.exit("No encontre '%s' en el resumen." % nombre)

ws.cell(row=fres("Reporting and exports"), column=7).value = 2
r = fres("Store file and location data")
ws.cell(row=r, column=6).value = 1
ws.cell(row=r, column=7).value = 1

wb.save(DESTINO)

# ---------- verificacion ----------
ws2 = openpyxl.load_workbook(DESTINO)["Support tracker"]
print("\nGuardado:", os.path.basename(DESTINO)); print(ws2["A2"].value, "\n")
tot = [0, 0, 0, 0]
for r in range(7, 18):
    for i, c in enumerate((4, 5, 6, 7)):
        tot[i] += ws2.cell(row=r, column=c).value or 0
print("resueltos %d | pend verif %d | esperando Ahold %d | activos %d | total %d"
      % (tot[0], tot[1], tot[2], tot[3], sum(tot)))
esp = {"RCS-28693": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28667": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28683": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28849": ("Accepted into Development", "Active with Appriss"),
       "RCS-28910": ("Work in progress",          "Active with Appriss"),
       "RCS-28824": ("Under Engineering Review",  "Active with Appriss"),
       "RCS-28944": ("Under Support Review",      "Active with Appriss"),
       "RCS-28589": ("Accepted into Development", "Active with Appriss"),
       "RCS-28911": ("Under Customer Review",     "Awaiting Ahold input"),
       "RCS-28281": ("Under Customer Review",     "Awaiting Ahold input")}
ok, vistos, claves = True, set(), []
print("\nAbiertos:")
for r in range(23, ws2.max_row + 1):
    k = ws2.cell(row=r, column=1).value
    if not (k and str(k).startswith("RCS-")):
        continue
    claves.append(k)
    cat = ws2.cell(row=r, column=4).value
    if cat == "Resolved":
        continue
    est = ws2.cell(row=r, column=3).value
    vistos.add(k)
    mal = k in esp and (est, cat) != esp[k]
    ok = ok and not mal
    print("  %s fila %-4d %-11s %-26s %s" % ("!!" if mal else "  ", r, k, est, cat))
if set(esp) - vistos:
    ok = False; print("\nFaltan:", ", ".join(sorted(set(esp) - vistos)))
if vistos - set(esp):
    ok = False; print("\nAbiertos inesperados:", ", ".join(sorted(vistos - set(esp))))
if sum(tot) != 55 or tot[0] != 45:
    ok = False; print("\nTotales mal, deberian ser 55 y 45 resueltos.")
d = sorted({k for k in claves if claves.count(k) > 1})
if d:
    ok = False; print("\nDuplicados:", ", ".join(d))
print("\n" + ("Todo cuadra." if ok else "Revisa lo marcado."))
